import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_services_docx(output_path):
    doc = docx.Document()

    # Page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("VISHWANATH SHARMA")
    run_title.font.name = 'Georgia'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Senior Corporate Financial & Investment Consultant")
    run_sub.font.name = 'Georgia'
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_contact = p_contact.add_run("Email: sharmavn2001@yahoo.com  |  Mobile: +91-9724305602  |  Location: Vadodara / Ahmedabad, Gujarat")
    run_contact.font.size = Pt(10)
    run_contact.font.italic = True
    run_contact.font.color.rgb = RGBColor(0x57, 0x53, 0x4E)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Overview Section
    p_heading = doc.add_paragraph()
    run_h = p_heading.add_run("CORPORATE FINANCIAL CONSULTANCY SERVICES")
    run_h.font.name = 'Georgia'
    run_h.font.size = Pt(14)
    run_h.font.bold = True
    run_h.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)
    
    p_intro = doc.add_paragraph(
        "With over 25 years of institutional banking executive experience across IndusInd Bank, Kotak Mahindra Bank, HDFC Bank, ICICI Bank, YES Bank, and ESAF Small Finance Bank, Vishwanath Sharma provides direct principal corporate finance consultancy to mid-market enterprises, CFOs, corporate promoters, and institutional lenders."
    )
    p_intro.paragraph_format.space_after = Pt(12)

    # 7 Practice Areas
    practice_areas = [
        {
            "title": "1. Debt Solutions & Structured Finance",
            "category": "Capital Raising",
            "desc": "Structuring and syndicating project finance, term loans, working capital facilities (CC/OD/LC/BG), and External Commercial Borrowings (ECB) for mid-market corporates.",
            "scope": [
                "Project Finance & Term Loans for Infrastructure, Energy, & Manufacturing",
                "Working Capital Facilities (Cash Credit, Overdraft, LC, Bank Guarantees)",
                "External Commercial Borrowings (ECB) & Cross-Border Buyer's Credit",
                "Equipment Finance, Factoring, & Asset-Backed Lending",
                "Leveraged Buyouts (LBOs) & Credit Rating Consultancy"
            ]
        },
        {
            "title": "2. Equity Consultancy & M&A Solutions",
            "category": "Capital Markets",
            "desc": "Advising growth-stage companies and corporate promoters on Private Equity (PE/VC) fundraising, buy-side & sell-side M&A transactions, takeovers, buybacks, and public market listings.",
            "scope": [
                "Private Equity (PE) & Venture Capital (VC) Growth Capital Raising",
                "Buy-side & Sell-side Mergers & Acquisitions (M&A) Execution",
                "Takeovers, Buyback Offers, & Stock Exchange De-Listing Consultancy",
                "IPO, FPO, Rights Issue, & Qualified Institutional Placements (QIP)",
                "Term Sheet Negotiation & Investor Readiness Structuring"
            ]
        },
        {
            "title": "3. Valuation & Fairness Opinions",
            "category": "Corporate Governance",
            "desc": "Independent, regulatory-compliant valuation reports for business equity, real estate, plant & machinery, ESOPs, and intangible assets under Income Tax Act, FEMA/RBI, and Companies Act.",
            "scope": [
                "Tangible Asset Valuation: Land & Building, Plant & Machinery",
                "Financial Assets & Securities Valuation (DCF, Multiples, Asset-Based)",
                "Regulatory Valuations under Income Tax Act, FEMA/RBI, & Companies Act",
                "ESOP Structuring, Fair Market Valuation, & Impairment Testing (Ind AS/IFRS)",
                "Brand & Intangible Asset Valuation"
            ]
        },
        {
            "title": "4. Techno-Economic Viability (TEV) & Feasibility Studies",
            "category": "Lender Consultancy",
            "desc": "Authoring bankable TEV study reports and conducting technical risk assessments required by consortium lenders prior to multi-crore credit sanctions.",
            "scope": [
                "Bankable Techno-Economic Viability (TEV) Study Reports",
                "Technical Risk & Manufacturing Capability Assessment",
                "Market Demand Evaluation & Financial Sensitivity Modeling",
                "Lender's Independent Engineer (LIE) Services & Drawdown Verification",
                "Detailed Project Report (DPR) Preparation for Bank Approvals"
            ]
        },
        {
            "title": "5. Stressed Assets & Insolvency (IBC Consultancy)",
            "category": "Turnaround Consultancy",
            "desc": "Specialized out-of-court debt restructuring, One-Time Settlement (OTS) structuring, and NCLT resolution plan formulation under the Insolvency and Bankruptcy Code (IBC).",
            "scope": [
                "Corporate Insolvency Resolution Process (CIRP) Consultancy",
                "Formulation & Submission of Bankable Resolution Plans",
                "Out-of-Court Debt Restructuring & One-Time Settlement (OTS)",
                "Promoter Handholding with Minimal Franchise & Financial Risk Exposure",
                "Liquidation Management & Stressed Asset Buying Consultancy"
            ]
        },
        {
            "title": "6. Corporate Training & Executive Masterclasses",
            "category": "Capability Building",
            "desc": "Executive masterclasses for banking credit teams and CFOs on paperless credit delivery, GSTR/ITR/Bank diagnostics, financial statement analysis, and IBC regulations.",
            "scope": [
                "Paperless Credit Delivery & Balance Sheet Analysis for Banking Staff",
                "GSTR, Income Tax Return, & Bank Statement Diagnostic Training",
                "Corporate Finance, Valuation Methods, & IBC Regulations Workshops",
                "Soft Skills, Leadership, & Team Building for Finance Leaders",
                "1-on-1 Executive Coaching for CFOs and Finance Controllers"
            ]
        },
        {
            "title": "7. Audit, Due Diligence & Financial Analytics",
            "category": "Assurance & Risk",
            "desc": "Buy-side and sell-side financial due diligence, forensic financial audits, fund flow verification, and integrated multi-source data reconciliations.",
            "scope": [
                "Buy-side & Sell-side Financial Due Diligence (FDD)",
                "Forensic Financial Audits & Fund Flow Verification",
                "Integrated Data Analytics (Bank Statement, GSTR, ITR Reconciliation)",
                "Internal Controls Assessment & Risk Management Audits",
                "Compliance Verification & KYC Validation for Institutional Lenders"
            ]
        }
    ]

    for area in practice_areas:
        p_sec = doc.add_paragraph()
        run_sec = p_sec.add_run(area["title"])
        run_sec.font.name = 'Georgia'
        run_sec.font.size = Pt(12)
        run_sec.font.bold = True
        run_sec.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

        p_desc = doc.add_paragraph(area["desc"])
        p_desc.paragraph_format.space_after = Pt(4)

        p_scope_head = doc.add_paragraph()
        run_sh = p_scope_head.add_run("Scope of Engagement:")
        run_sh.font.bold = True
        run_sh.font.size = Pt(10)
        run_sh.font.color.rgb = RGBColor(0x57, 0x53, 0x4E)

        for item in area["scope"]:
            p_item = doc.add_paragraph(style='List Bullet')
            p_item.paragraph_format.space_after = Pt(2)
            run_item = p_item.add_run(item)
            run_item.font.size = Pt(10)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Footer note
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ft = p_footer.add_run("Confidential Corporate Financial Consultancy Document  |  © Vishwanath Sharma")
    run_ft.font.size = Pt(9)
    run_ft.font.italic = True
    run_ft.font.color.rgb = RGBColor(0x78, 0x71, 0x6C)

    doc.save(output_path)
    print(f"Created clean Word document at {output_path}")

if __name__ == "__main__":
    os.makedirs("public/service_pdfs", exist_ok=True)
    os.makedirs("service_pdfs", exist_ok=True)
    create_services_docx("public/service_pdfs/Vishwanath_Sharma_Corporate_Consultancy_Services.docx")
    create_services_docx("service_pdfs/Vishwanath_Sharma_Corporate_Consultancy_Services.docx")
